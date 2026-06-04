# -*- coding: utf-8 -*-
from pathlib import Path

from docx import Document


docx = list(Path(r"C:\coding\YOLO-FD").glob("*修改版.docx"))[0]
doc = Document(docx)
print("doc", docx.name)
print("paragraphs", len(doc.paragraphs), "tables", len(doc.tables), "inline_shapes", len(doc.inline_shapes))

needles = [
    "图1 YOLO-FD 两阶段训练与推理校准算法框架",
    "图2 IoU 计算公式",
    "图3 YOLO-FD 检测分支损失公式",
    "图4 不确定性多任务损失公式",
    "图5 PCGrad 梯度冲突修正公式",
    "图6 论文 SOTA、完整验证集与测试子集指标对比",
    "图7 第二阶段训练曲线",
    "图8 完整验证集 PR 曲线",
    "图9 完整验证集 F1 曲线",
    "图10 完整验证集归一化混淆矩阵",
    "图11 测试子集部分预测可视化结果",
    "本文没有改变 YOLO-FD 的主体架构",
    "mAP50 从 0.942 提升到 0.947",
]
full_text = "\n".join(p.text for p in doc.paragraphs)
for needle in needles:
    print(needle, needle in full_text)

print("\nselected paragraphs:")
for index, paragraph in enumerate(doc.paragraphs):
    text = " ".join(paragraph.text.split())
    if text.startswith(("图", "3.", "4.")) or "本文没有改变 YOLO-FD" in text or "mAP50 从" in text:
        print(f"{index:04d} {text[:180]}")
